"""Mentor dashboard endpoints — review unanswered questions."""
import io
import math
from typing import Dict, List, Tuple

import structlog
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form, status
from sqlalchemy import select, desc, func, update
from sqlalchemy.ext.asyncio import AsyncSession

from models.database import get_db, UnansweredPool, QAPair
from models.schemas import (
    UnansweredPoolResponse,
    MentorAnswerRequest,
    DismissRequest,
    UnansweredGroupResponse,
    UnansweredVariant,
)
from api.middleware.auth import require_api_key

log = structlog.get_logger()
router = APIRouter(prefix="/mentor", tags=["mentor"])

# Cosine similarity threshold for "same cluster". Tuned for text-embedding-3-small:
# 0.82 groups paraphrases ("how do I start a startup", "tips on founding a company")
# without dragging in loosely-related items.
CLUSTER_SIMILARITY_THRESHOLD = 0.82

# Process-level cache of question text → embedding vector so repeated GETs on
# /mentor/unanswered/grouped don't re-embed the same rows over and over.
_embedding_cache: Dict[str, List[float]] = {}


def _cosine(a: List[float], b: List[float]) -> float:
    if not a or not b:
        return 0.0
    dot = 0.0
    na = 0.0
    nb = 0.0
    for x, y in zip(a, b):
        dot += x * y
        na += x * x
        nb += y * y
    denom = math.sqrt(na) * math.sqrt(nb)
    return dot / denom if denom else 0.0


@router.get("/unanswered", response_model=list[UnansweredPoolResponse])
async def list_unanswered(
    status_filter: str = "pending",
    limit: int = 50,
    db: AsyncSession = Depends(get_db),
    _: str = Depends(require_api_key),
):
    """List questions in the unanswered pool for mentor review (raw, ungrouped)."""
    query = (
        select(UnansweredPool)
        .where(UnansweredPool.status == status_filter)
        .order_by(desc(UnansweredPool.created_at))
        .limit(limit)
    )
    result = await db.execute(query)
    return result.scalars().all()


async def _ensure_clusters(
    db: AsyncSession, rows: list[UnansweredPool]
) -> dict[int, list[float]]:
    """Assign `cluster_id` to any row that doesn't have one via greedy
    embedding-based clustering.

    Returns a mapping `row_id -> embedding_vector` for use by the caller
    (e.g. to compute per-variant similarity to the group representative).

    Strategy:
      1. Load or compute embeddings for all supplied rows (cached in-process).
      2. Seed the cluster map with rows that already have a `cluster_id`.
      3. For each unassigned row, compare against the centroid-like vector of
         each existing cluster (we use the first row's vector of the cluster
         as a cheap proxy). Join if cosine >= threshold, else allocate a new
         cluster_id.
      4. Persist the new assignments.
    """
    import app_state as _app_state
    embedding_svc = _app_state.state.get("embedding")
    if embedding_svc is None:
        # No embedding service configured — fall back to one cluster per row.
        next_id = 1
        updated: list[int] = []
        for row in rows:
            if row.cluster_id is None:
                row.cluster_id = next_id
                next_id += 1
                updated.append(row.id)
        if updated:
            await db.flush()
        return {}

    # ── Step 1: embeddings for every row ─────────────────────────────────
    row_vecs: dict[int, list[float]] = {}
    for row in rows:
        key = row.question.strip().lower()
        vec = _embedding_cache.get(key)
        if vec is None:
            try:
                vec = await embedding_svc.embed(row.question)
            except Exception as e:
                log.warning("embed_failed_for_cluster", id=row.id, error=str(e))
                continue
            _embedding_cache[key] = vec
        row_vecs[row.id] = vec

    # ── Step 2: seed existing clusters (pick one representative vector each)
    cluster_rep_vec: dict[int, list[float]] = {}
    for row in rows:
        if row.cluster_id is not None and row.cluster_id not in cluster_rep_vec:
            v = row_vecs.get(row.id)
            if v is not None:
                cluster_rep_vec[row.cluster_id] = v

    # Next available cluster id — 1 + max(existing)
    next_cluster_id = (max(cluster_rep_vec.keys(), default=0) or 0) + 1

    # ── Step 3: assign unassigned rows greedily ──────────────────────────
    assigned = False
    for row in rows:
        if row.cluster_id is not None:
            continue
        vec = row_vecs.get(row.id)
        if vec is None:
            row.cluster_id = next_cluster_id
            next_cluster_id += 1
            assigned = True
            continue

        best_cid = None
        best_sim = -1.0
        for cid, cvec in cluster_rep_vec.items():
            sim = _cosine(vec, cvec)
            if sim > best_sim:
                best_sim = sim
                best_cid = cid

        if best_cid is not None and best_sim >= CLUSTER_SIMILARITY_THRESHOLD:
            row.cluster_id = best_cid
        else:
            row.cluster_id = next_cluster_id
            cluster_rep_vec[next_cluster_id] = vec
            next_cluster_id += 1
        assigned = True

    if assigned:
        await db.flush()

    return row_vecs


@router.get("/unanswered/grouped", response_model=list[UnansweredGroupResponse])
async def list_unanswered_grouped(
    status_filter: str = "pending",
    limit: int = 50,
    db: AsyncSession = Depends(get_db),
    _: str = Depends(require_api_key),
):
    """List unanswered questions grouped by semantic similarity.

    The first time a row is seen here its `cluster_id` is assigned via greedy
    embedding clustering and persisted, so subsequent calls are fast and
    mentor-driven splits (see `/unanswered/{id}/split`) stick.
    """
    all_rows = (
        await db.execute(
            select(UnansweredPool)
            .where(UnansweredPool.status == status_filter)
            .order_by(desc(UnansweredPool.created_at))
        )
    ).scalars().all()
    if not all_rows:
        return []

    row_vecs = await _ensure_clusters(db, list(all_rows))

    # Group in python by cluster_id
    by_cluster: dict[int, list[UnansweredPool]] = {}
    for row in all_rows:
        by_cluster.setdefault(row.cluster_id or -row.id, []).append(row)

    groups: list[UnansweredGroupResponse] = []
    for cid, members in by_cluster.items():
        # Representative = most recently created row in the cluster
        members_sorted = sorted(members, key=lambda r: r.created_at, reverse=True)
        rep = members_sorted[0]

        rep_vec = row_vecs.get(rep.id)
        variants: list[UnansweredVariant] = []
        for m in members_sorted:
            sim = 1.0
            if rep_vec is not None:
                mv = row_vecs.get(m.id)
                if mv is not None and m.id != rep.id:
                    sim = round(_cosine(rep_vec, mv), 4)
            variants.append(
                UnansweredVariant(
                    id=m.id,
                    question=m.question,
                    similarity=sim,
                    created_at=m.created_at,
                )
            )

        groups.append(
            UnansweredGroupResponse(
                cluster_id=int(cid) if cid >= 0 else -1,
                representative_id=rep.id,
                question=rep.question,
                count=len(members),
                group_ids=[m.id for m in members],
                variants=variants,
                general_response=rep.general_response,
                rag_context_used=rep.rag_context_used,
                confidence_score=rep.confidence_score,
                mentor_id=rep.mentor_id,
                created_at=min(m.created_at for m in members),
                latest_created_at=max(m.created_at for m in members),
            )
        )

    # Sort groups by most recent activity, then trim to limit
    groups.sort(key=lambda g: g.latest_created_at, reverse=True)
    return groups[:limit]


@router.post("/unanswered/{item_id}/split", response_model=UnansweredPoolResponse)
async def split_from_cluster(
    item_id: int,
    db: AsyncSession = Depends(get_db),
    _: str = Depends(require_api_key),
):
    """Move one row out of its current cluster into a brand-new singleton
    cluster. Use this when the mentor judges a grouped question as "not
    actually similar" to the others.
    """
    row = (
        await db.execute(select(UnansweredPool).where(UnansweredPool.id == item_id))
    ).scalar_one_or_none()
    if not row:
        raise HTTPException(status_code=404, detail="Unanswered item not found")

    max_cid = (
        await db.execute(select(func.max(UnansweredPool.cluster_id)))
    ).scalar() or 0
    row.cluster_id = int(max_cid) + 1
    await db.flush()
    await db.refresh(row)
    log.info("cluster_split", id=item_id, new_cluster_id=row.cluster_id)
    return row


@router.post("/answer/{item_id}", response_model=UnansweredPoolResponse)
async def answer_question(
    item_id: int,
    payload: MentorAnswerRequest,
    db: AsyncSession = Depends(get_db),
    _: str = Depends(require_api_key),
):
    """Mentor provides an authoritative answer for an unanswered question.

    If `group_ids` is provided, every row in the group is marked answered with
    the same mentor_answer (the KB row is still created only once).
    """
    from datetime import datetime

    result = await db.execute(select(UnansweredPool).where(UnansweredPool.id == item_id))
    item = result.scalar_one_or_none()
    if not item:
        raise HTTPException(status_code=404, detail="Unanswered item not found")

    # Targets: representative + any duplicates the caller grouped with it
    target_ids = set(payload.group_ids or [])
    target_ids.add(item_id)

    now = datetime.utcnow()
    targets = (
        await db.execute(select(UnansweredPool).where(UnansweredPool.id.in_(target_ids)))
    ).scalars().all()

    # Detect verbatim copy-paste of the AI RAG response. If the mentor pasted
    # the AI answer unchanged, this row should NOT influence threshold
    # recommendations (the AI was already right — we don't want to claim "we
    # should have hit KB here"). Whitespace-insensitive comparison.
    def _norm(s: str | None) -> str:
        return " ".join((s or "").split()).strip().lower()

    mentor_norm = _norm(payload.mentor_answer)
    for row in targets:
        row.mentor_answer = payload.mentor_answer
        row.status = "answered"
        row.reviewed_at = now
        row.answer_was_copied = bool(
            row.general_response and mentor_norm == _norm(row.general_response)
        )
    await db.flush()

    # Optionally add to knowledge base — only one KB row for the whole group
    if payload.add_to_kb:
        qa = QAPair(
            question=item.question,
            answer=payload.mentor_answer,
            mentor_id=payload.mentor_id,
            source="session",
            approved=True,
        )
        db.add(qa)
        await db.flush()
        await db.refresh(qa)

        from api.routes.knowledge import _index_qa
        await _index_qa(qa)
        log.info(
            "mentor_answer_added_to_kb",
            qa_id=qa.id,
            group_size=len(targets),
        )

    await db.refresh(item)
    return item


@router.get("/threshold/recommend")
async def recommend_threshold(
    mentor_id: str | None = None,
    db: AsyncSession = Depends(get_db),
    _: str = Depends(require_api_key),
):
    """Suggest a similarity threshold based on the distribution of routing
    scores for Case A questions the mentor has answered with original text.

    Rationale
    ---------
    Every Case A row stores the best cosine similarity against approved_qa at
    routing time (`confidence_score`). If a row was answered by the mentor
    with *original* text (not a verbatim copy of the AI RAG response), that
    question was genuinely missing from the KB — the router correctly routed
    it away from Case B. Rows where the mentor copy-pasted the AI response
    are excluded: they don't tell us the threshold was wrong, they tell us
    the AI was good enough.

    From the remaining distribution we take a high percentile (P90) of
    confidence scores, bumped up by 0.01 as a safety buffer. Intuition: 90%
    of genuinely-original mentor answers had similarity below this number,
    so a threshold here would have correctly kept them in Case A while still
    giving every truly-KB-matching question a chance to hit Case B.

    This endpoint **never** mutates config — it only reports. The mentor
    applies the change manually via SIMILARITY_THRESHOLD in .env.
    """
    from config import get_settings
    settings = get_settings()
    current = settings.SIMILARITY_THRESHOLD

    q = select(UnansweredPool).where(UnansweredPool.status == "answered")
    if mentor_id:
        q = q.where(UnansweredPool.mentor_id == mentor_id)
    rows = (await db.execute(q)).scalars().all()

    all_scored = [r for r in rows if r.confidence_score is not None]
    eligible = [r for r in all_scored if not r.answer_was_copied]
    excluded_copied = len(all_scored) - len(eligible)

    def pct(values: list[float], p: float) -> float:
        if not values:
            return 0.0
        s = sorted(values)
        k = max(0, min(len(s) - 1, int(round((p / 100.0) * (len(s) - 1)))))
        return s[k]

    # Bucket counts for the histogram shown in the UI
    buckets = [(round(x / 20, 2), round((x + 1) / 20, 2)) for x in range(20)]  # 0.00–1.00 in 0.05 steps
    histogram: list[dict] = []
    eligible_scores = [float(r.confidence_score) for r in eligible]
    for lo, hi in buckets:
        count = sum(1 for s in eligible_scores if lo <= s < hi or (hi == 1.0 and s == 1.0))
        histogram.append({"lo": lo, "hi": hi, "count": count})

    if len(eligible_scores) < 5:
        return {
            "current_threshold": current,
            "suggested_threshold": current,
            "sample_size": len(eligible_scores),
            "excluded_copied": excluded_copied,
            "percentiles": {},
            "histogram": histogram,
            "reasoning": (
                f"Not enough data yet ({len(eligible_scores)} mentor-original answers). "
                "Need at least 5 to make a recommendation. Keep answering unanswered "
                "questions without copy-pasting the AI response."
            ),
            "mentor_id": mentor_id,
        }

    p50 = round(pct(eligible_scores, 50), 4)
    p75 = round(pct(eligible_scores, 75), 4)
    p90 = round(pct(eligible_scores, 90), 4)
    p95 = round(pct(eligible_scores, 95), 4)
    max_s = round(max(eligible_scores), 4)

    # Safety buffer so a borderline sample doesn't put us right at the edge
    suggested = round(min(0.98, max(0.5, p90 + 0.01)), 4)

    reasoning = (
        f"Analyzed {len(eligible_scores)} answered Case A rows where the mentor "
        f"wrote an original answer (excluded {excluded_copied} copy-pasted ones). "
        f"Their similarity-to-KB scores spread from {round(min(eligible_scores),4)} "
        f"to {max_s}. Using the 90th percentile ({p90}) + 0.01 safety buffer = "
        f"{suggested}. A threshold at this level would have kept 90% of these "
        f"genuinely-new questions in Case A (where they belonged), while "
        f"still letting near-duplicate questions hit Case B."
    )

    return {
        "current_threshold": current,
        "suggested_threshold": suggested,
        "sample_size": len(eligible_scores),
        "excluded_copied": excluded_copied,
        "percentiles": {"p50": p50, "p75": p75, "p90": p90, "p95": p95, "max": max_s},
        "histogram": histogram,
        "reasoning": reasoning,
        "mentor_id": mentor_id,
    }


@router.post("/dismiss/{item_id}", response_model=UnansweredPoolResponse)
async def dismiss_question(
    item_id: int,
    payload: DismissRequest | None = None,
    db: AsyncSession = Depends(get_db),
    _: str = Depends(require_api_key),
):
    """Dismiss a question (spam, off-topic, etc.).

    If `group_ids` is supplied in the body, dismiss every id in one call.
    """
    result = await db.execute(select(UnansweredPool).where(UnansweredPool.id == item_id))
    item = result.scalar_one_or_none()
    if not item:
        raise HTTPException(status_code=404, detail="Item not found")

    target_ids = set((payload.group_ids if payload else None) or [])
    target_ids.add(item_id)
    targets = (
        await db.execute(select(UnansweredPool).where(UnansweredPool.id.in_(target_ids)))
    ).scalars().all()
    for row in targets:
        row.status = "dismissed"
    await db.flush()
    await db.refresh(item)
    return item


def _chunk_text_inline(text: str, chunk_size: int = 400, overlap: int = 80) -> list[str]:
    """Inlined chunker — no import from workers.transcript_processor."""
    import re
    text = re.sub(r'\s+', ' ', text.strip())
    chunks: list[str] = []
    start = 0
    while start < len(text):
        end = min(start + chunk_size, len(text))
        if end < len(text):
            for sep in ['. ', '! ', '? ', '\n']:
                pos = text.rfind(sep, start + overlap, end)
                if pos != -1:
                    end = pos + len(sep)
                    break
        chunks.append(text[start:end].strip())
        if end >= len(text):
            break
        start = end - overlap
    return [c for c in chunks if len(c) > 20]


@router.post("/transcript", status_code=status.HTTP_202_ACCEPTED)
async def ingest_transcript(
    payload: dict,
    _: str = Depends(require_api_key),
):
    """Submit transcript — chunk + embed into session_chunks (in-process dev)."""
    log.info("transcript_route_entered", payload_keys=list(payload.keys()))
    import hashlib

    transcript_text = payload.get("transcript_text", "").strip()
    mentor_id = payload.get("mentor_id", "default")
    source = payload.get("source", "session")
    title = (payload.get("title") or "").strip()

    if not transcript_text:
        raise HTTPException(status_code=400, detail="transcript_text required")

    # Auto-name untitled transcripts so RAG sources still have something
    # meaningful to display in the mentor dashboard.
    if not title:
        from datetime import datetime
        title = f"Untitled {source} · {datetime.utcnow().strftime('%Y-%m-%d %H:%M')}"

    log.info("transcript_payload_ok", length=len(transcript_text), mentor_id=mentor_id, title=title)

    import app_state as _app_state
    embedding_svc = _app_state.state.get("embedding")
    vector_store = _app_state.state.get("vector_store")
    if embedding_svc is None or vector_store is None:
        log.error("transcript_services_missing",
                  has_embedding=embedding_svc is not None,
                  has_vector_store=vector_store is not None)
        raise HTTPException(status_code=503, detail="embedding/vector_store not initialized")

    chunks = _chunk_text_inline(transcript_text)
    log.info("transcript_chunked", chunks=len(chunks), mentor_id=mentor_id)

    for i, chunk in enumerate(chunks):
        log.info("transcript_embedding_chunk", i=i, len=len(chunk))
        vector = await embedding_svc.embed(chunk)
        chunk_id = int(hashlib.md5(chunk.encode()).hexdigest()[:16], 16) % (2**63)
        await vector_store.upsert(
            collection="session_chunks",
            vector_id=chunk_id,
            vector=vector,
            payload={
                "text": chunk,
                "mentor_id": mentor_id,
                "source": source,
                "transcript_title": title,
                "chunk_index": i,
            },
        )
        log.info("transcript_chunk_upserted", i=i, chunk_id=chunk_id)

    log.info("transcript_done", chunks=len(chunks))
    return {"task_id": f"inproc-{mentor_id}", "status": "done", "mode": "inproc", "chunks": len(chunks)}


# ── File-based transcript upload ──────────────────────────────────────────
SUPPORTED_EXTS = {".txt", ".md", ".pdf", ".docx", ".pptx"}
MAX_UPLOAD_BYTES = 25 * 1024 * 1024  # 25 MB


def _extract_text_from_file(filename: str, data: bytes) -> str:
    """Extract plain text from txt / md / pdf / docx / pptx bytes."""
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in SUPPORTED_EXTS:
        raise HTTPException(
            status_code=415,
            detail=f"Unsupported file type {ext!r}. Supported: {sorted(SUPPORTED_EXTS)}",
        )

    if ext in (".txt", ".md"):
        for enc in ("utf-8", "utf-16", "latin-1"):
            try:
                return data.decode(enc)
            except UnicodeDecodeError:
                continue
        raise HTTPException(status_code=400, detail="Could not decode text file")

    if ext == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError:
            raise HTTPException(status_code=500, detail="pypdf not installed")
        reader = PdfReader(io.BytesIO(data))
        return "\n\n".join((p.extract_text() or "") for p in reader.pages).strip()

    if ext == ".docx":
        try:
            from docx import Document
        except ImportError:
            raise HTTPException(status_code=500, detail="python-docx not installed")
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts).strip()

    if ext == ".pptx":
        try:
            from pptx import Presentation
        except ImportError:
            raise HTTPException(status_code=500, detail="python-pptx not installed")
        prs = Presentation(io.BytesIO(data))
        slides = []
        for i, slide in enumerate(prs.slides, 1):
            slide_lines = [f"[Slide {i}]"]
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        text = "".join(r.text for r in para.runs).strip()
                        if text:
                            slide_lines.append(text)
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        cells = [c.text.strip() for c in row.cells if c.text.strip()]
                        if cells:
                            slide_lines.append(" | ".join(cells))
            # Speaker notes
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    slide_lines.append(f"(notes) {notes}")
            slides.append("\n".join(slide_lines))
        return "\n\n".join(slides).strip()

    raise HTTPException(status_code=415, detail=f"Unsupported file type {ext!r}")


@router.post("/transcript/upload", status_code=status.HTTP_202_ACCEPTED)
async def ingest_transcript_file(
    file: UploadFile = File(...),
    mentor_id: str = Form("default"),
    source: str = Form("session"),
    title: str = Form(""),
    _: str = Depends(require_api_key),
):
    """Upload txt / md / pdf / docx / pptx → extract → chunk + embed."""
    import hashlib
    from datetime import datetime

    raw = await file.read()
    if not raw:
        raise HTTPException(status_code=400, detail="empty file")
    if len(raw) > MAX_UPLOAD_BYTES:
        raise HTTPException(
            status_code=413,
            detail=f"File too large ({len(raw)} bytes > {MAX_UPLOAD_BYTES})",
        )

    text = _extract_text_from_file(file.filename or "upload.txt", raw).strip()
    if not text:
        raise HTTPException(status_code=400, detail="No extractable text found in file")

    title = (title or "").strip() or f"{file.filename} · {datetime.utcnow().strftime('%Y-%m-%d %H:%M')}"

    log.info(
        "transcript_upload_extracted",
        filename=file.filename,
        bytes=len(raw),
        chars=len(text),
        mentor_id=mentor_id,
    )

    import app_state as _app_state
    embedding_svc = _app_state.state.get("embedding")
    vector_store = _app_state.state.get("vector_store")
    if embedding_svc is None or vector_store is None:
        raise HTTPException(status_code=503, detail="embedding/vector_store not initialized")

    chunks = _chunk_text_inline(text)
    for i, chunk in enumerate(chunks):
        vector = await embedding_svc.embed(chunk)
        chunk_id = int(hashlib.md5(chunk.encode()).hexdigest()[:16], 16) % (2**63)
        await vector_store.upsert(
            collection="session_chunks",
            vector_id=chunk_id,
            vector=vector,
            payload={
                "text": chunk,
                "mentor_id": mentor_id,
                "source": source,
                "transcript_title": title,
                "chunk_index": i,
                "filename": file.filename,
            },
        )

    log.info("transcript_upload_done", filename=file.filename, chunks=len(chunks))
    return {
        "task_id": f"inproc-{mentor_id}",
        "status": "done",
        "mode": "inproc",
        "chunks": len(chunks),
        "chars": len(text),
        "title": title,
        "filename": file.filename,
    }


# ── RAG inspection / management ────────────────────────────────────────────
@router.get("/rag/sources")
async def list_rag_sources(_: str = Depends(require_api_key)):
    """List every distinct (transcript_title, source, mentor_id, filename)
    indexed in session_chunks, with a chunk count per source."""
    import app_state as _app_state
    vector_store = _app_state.state.get("vector_store")
    if vector_store is None:
        raise HTTPException(status_code=503, detail="vector_store not initialized")

    points = await vector_store.scroll_all("session_chunks", limit=10_000)
    groups: Dict[Tuple[str, str, str, str], Dict] = {}
    for p in points:
        pl = p["payload"]
        key = (
            pl.get("transcript_title") or "(untitled)",
            pl.get("source") or "session",
            pl.get("mentor_id") or "default",
            pl.get("filename") or "",
        )
        g = groups.setdefault(key, {
            "title": key[0],
            "source": key[1],
            "mentor_id": key[2],
            "filename": key[3] or None,
            "chunk_count": 0,
            "first_chunk_id": p["id"],
        })
        g["chunk_count"] += 1

    sources = sorted(groups.values(), key=lambda g: -g["chunk_count"])
    return {
        "total_chunks": len(points),
        "total_sources": len(sources),
        "sources": sources,
    }


@router.get("/rag/chunks")
async def list_rag_chunks(
    title: str,
    _: str = Depends(require_api_key),
):
    """List every chunk for one transcript title (sorted by chunk_index)."""
    import app_state as _app_state
    vector_store = _app_state.state.get("vector_store")
    if vector_store is None:
        raise HTTPException(status_code=503, detail="vector_store not initialized")

    points = await vector_store.scroll_all(
        "session_chunks",
        filter_payload={"transcript_title": title},
        limit=5_000,
    )
    chunks = [
        {
            "id": p["id"],
            "text": p["payload"].get("text", ""),
            "title": p["payload"].get("transcript_title", ""),
            "source": p["payload"].get("source", ""),
            "mentor_id": p["payload"].get("mentor_id", ""),
            "chunk_index": p["payload"].get("chunk_index", 0),
            "filename": p["payload"].get("filename"),
        }
        for p in points
    ]
    chunks.sort(key=lambda c: c["chunk_index"])
    return chunks


@router.delete("/rag/sources")
async def delete_rag_source(
    title: str,
    _: str = Depends(require_api_key),
):
    """Remove every chunk for one transcript title."""
    import app_state as _app_state
    vector_store = _app_state.state.get("vector_store")
    if vector_store is None:
        raise HTTPException(status_code=503, detail="vector_store not initialized")

    deleted = await vector_store.delete_by_filter(
        "session_chunks", {"transcript_title": title}
    )
    log.info("rag_source_deleted", title=title, deleted=deleted)
    return {"deleted": deleted, "title": title}


@router.delete("/rag/chunks/{chunk_id}")
async def delete_rag_chunk(
    chunk_id: int,
    _: str = Depends(require_api_key),
):
    """Remove a single chunk by Qdrant point id."""
    import app_state as _app_state
    vector_store = _app_state.state.get("vector_store")
    if vector_store is None:
        raise HTTPException(status_code=503, detail="vector_store not initialized")

    await vector_store.delete("session_chunks", chunk_id)
    log.info("rag_chunk_deleted", id=chunk_id)
    return {"deleted": 1}


@router.delete("/rag/all")
async def delete_all_rag(_: str = Depends(require_api_key)):
    """Wipe the entire session_chunks collection (irreversible)."""
    import app_state as _app_state
    vector_store = _app_state.state.get("vector_store")
    if vector_store is None:
        raise HTTPException(status_code=503, detail="vector_store not initialized")

    deleted = await vector_store.delete_all("session_chunks")
    log.info("rag_all_deleted", deleted=deleted)
    return {"deleted_chunks": deleted}
